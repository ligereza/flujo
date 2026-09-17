#!/usr/bin/env python3
"""Exporta la tabla de testeos RD ya ordenada, en xlsx y en csv.

Es el tercer entregable del corpus: no un grafico ni un panel, sino **la
planilla que la planilla heredada deberia haber sido**. Una fila por muestra
analizada, con su origen exacto -- hoja y numero de fila -- para que cualquiera
pueda ir a verla en el archivo original y comprobarla.

Trae dos hojas:

- `muestras`   una fila por muestra, con los cuatro reactivos en columnas fijas
- `descartes`  lo que NO entro y por que, con su origen. Una cifra que descarta
               filas tiene que poder decir cuales.

Que NO hace: no dice que sustancia hay. `sustancia` es lo que la persona
declaro al entregar la muestra, y `color` es lo que se vio. El veredicto lo da
quien esta en la mesa, de viva voz, y no se registra.

    py tools/exportar_tabla_ordenada.py [--out <carpeta>]
"""
from __future__ import annotations

import argparse
import csv
import sqlite3
import sys
from pathlib import Path

_REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_REPO / "tools"))
sys.path.insert(0, str(_REPO / "src"))

from gen_resumen_testeos import ETIQUETA_COLOR, ETIQUETA_SUSTANCIA, reunir  # noqa: E402

from flujo.rd.paths import rd_db_path  # noqa: E402

COLUMNAS = [
    "periodo", "fecha", "fiesta", "hoja_origen", "fila_origen",
    "sustancia_declarada", "sustancia_tal_como_se_escribio", "formato",
    "reactivo_1", "color_1", "reactivo_2", "color_2",
    "reactivo_3", "color_3", "reactivo_4", "color_4",
    "reactivo_sin_nombrar",
]


def _filas(datos: dict) -> list[list]:
    et_s = ETIQUETA_SUSTANCIA
    et_c = ETIQUETA_COLOR
    salida = []
    for m in datos["tabla"]:
        fila = [
            m["periodo"], m["fecha"] or "", m["fiesta"], m["hoja"], m["fila"],
            et_s.get(m["sustancia"], m["sustancia"].lower()),
            m["sustancia_cruda"] or "", m["formato"] or "",
        ]
        for i in range(4):
            e = m["ensayos"][i] if i < len(m["ensayos"]) else None
            if e is None:
                fila += ["", ""]
            else:
                color = e["color"] or ""
                fila += [
                    (e["reactivo"] or "").replace("_", ":").lower() or "(sin nombrar)",
                    et_c.get(color, color.replace("_", " ").replace("+", " + ").lower()),
                ]
        fila.append("si" if m["reactivo_sin_nombrar"] else "")
        salida.append(fila)
    return salida


def _descartes(db: Path) -> list[list]:
    """Lo que quedo fuera, con su origen. Sin esto la tabla no se puede auditar."""
    conn = sqlite3.connect(f"file:{db.resolve().as_posix()}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    try:
        filas = []
        for r in conn.execute(
            "SELECT e.source_period_label p, e.source_sheet_name h, f.source_row n, "
            "f.substance_raw s, f.format_raw fo, f.row_duplicate_status d, "
            "f.copied_from_sheet o FROM testeo_filas_fuente f "
            "JOIN testeo_eventos_fuente e ON e.event_id = f.event_id "
            "WHERE f.row_duplicate_status <> 'first_occurrence' "
            "ORDER BY e.source_sheet_name, f.source_row"
        ):
            motivo = ("repetida dentro de su hoja"
                      if r["d"] == "repeat_within_sheet"
                      else f"copiada de «{r['o']}»" if r["o"] else "copiada de otra jornada")
            filas.append([r["p"], r["h"].strip(), r["n"],
                          r["s"] or "", r["fo"] or "", motivo])
        return filas
    finally:
        conn.close()


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--db", default=None)
    ap.add_argument("--out", default=str(_REPO / "dist_compartir"))
    args = ap.parse_args()

    db = Path(args.db) if args.db else rd_db_path(_REPO)
    if not db.is_file():
        print(f"no existe la base: {db}")
        return 1

    datos = reunir(db, detalle=True)
    filas = _filas(datos)
    descartes = _descartes(db)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    csv_path = out / "testeos_rd_ordenado.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNAS)
        w.writerows(filas)

    xlsx_path = None
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        pass
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "muestras"
        ws.append(COLUMNAS)
        for f in filas:
            ws.append(f)
        ws2 = wb.create_sheet("descartes")
        ws2.append(["periodo", "hoja_origen", "fila_origen", "sustancia",
                    "formato", "motivo"])
        for f in descartes:
            ws2.append(f)
        cabecera = Font(bold=True, color="FFEFE9DD")
        fondo = PatternFill("solid", fgColor="FF3B1A4A")
        for hoja in (ws, ws2):
            hoja.freeze_panes = "A2"
            for celda in hoja[1]:
                celda.font = cabecera
                celda.fill = fondo
                celda.alignment = Alignment(horizontal="center")
            for col in hoja.columns:
                largo = max(len(str(c.value or "")) for c in col[:400])
                hoja.column_dimensions[col[0].column_letter].width = min(max(largo + 2, 9), 34)
        xlsx_path = out / "testeos_rd_ordenado.xlsx"
        wb.save(xlsx_path)

    t = datos["totales"]
    print(f"{csv_path}")
    if xlsx_path:
        print(f"{xlsx_path}")
    print(f"  {len(filas)} muestras · {t['jornadas']} fiestas · "
          f"{len(descartes)} filas descartadas con su motivo")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
